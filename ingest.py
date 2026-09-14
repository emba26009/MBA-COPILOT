import os, zipfile, tempfile
from pathlib import Path
import db
from PyPDF2 import PdfReader
from docx import Document
from pptx import Presentation
import pandas as pd

SUPPORTED={'.pdf','.docx','.pptx','.xlsx','.xls'}
MAX_ZIP_DEPTH=3

def clean(s): return ' '.join(str(s or '').split()).strip()
def chunk(text,size=1800,overlap=250):
    text=clean(text)
    if not text:return []
    out=[];i=0
    while i<len(text):
        j=min(len(text),i+size);part=text[i:j]
        if len(part)>80:out.append(part)
        if j==len(text):break
        i=max(i+size-overlap,i+1)
    return out

def extract(path):
    ext=path.suffix.lower();out=[]
    if ext=='.pdf':
        for n,p in enumerate(PdfReader(str(path)).pages,1):
            parts=chunk(p.extract_text())
            for k,c in enumerate(parts):out.append((f'Page {n}'+(f' · Part {k+1}' if len(parts)>1 else ''),c))
    elif ext=='.docx':
        parts=chunk('\n'.join(p.text for p in Document(str(path)).paragraphs))
        for k,c in enumerate(parts):out.append((f'Paragraphs · Part {k+1}',c))
    elif ext=='.pptx':
        for n,slide in enumerate(Presentation(str(path)).slides,1):
            parts=chunk(' '.join(sh.text for sh in slide.shapes if hasattr(sh,'text')))
            for k,c in enumerate(parts):out.append((f'Slide {n}'+(f' · Part {k+1}' if len(parts)>1 else ''),c))
    elif ext in ('.xlsx','.xls'):
        book=pd.ExcelFile(str(path))
        for sheet in book.sheet_names:
            df=pd.read_excel(str(path),sheet_name=sheet,header=None).fillna('');rows=[]
            for _,r in df.iterrows():
                vals=[clean(x) for x in r.tolist() if clean(x)]
                if vals:rows.append(' | '.join(vals))
            parts=chunk('\n'.join(rows))
            for k,c in enumerate(parts):out.append((f'Sheet {sheet}'+(f' · Part {k+1}' if len(parts)>1 else ''),c))
    return out

def embed_and_store(filename,subject,path):
    passages=extract(path)
    if not passages:return {'filename':filename,'passages':0,'embeddings':0}
    did=db.add_document(filename,subject,path.suffix.lower(),passages)
    embedded=0
    if os.getenv('OPENAI_API_KEY'):
        try:
            from openai import OpenAI
            client=OpenAI(api_key=os.environ['OPENAI_API_KEY']);vectors=[]
            for start in range(0,len(passages),64):
                r=client.embeddings.create(model='text-embedding-3-small',input=[x[1] for x in passages[start:start+64]])
                vectors.extend(x.embedding for x in r.data)
            db.store_embeddings(did,vectors);embedded=len(vectors)
        except Exception as e: print('Embedding warning:',e)
    return {'filename':filename,'passages':len(passages),'embeddings':embedded}

def _safe_extract(z,root):
    root=Path(root).resolve()
    for member in z.infolist():
        target=(root/member.filename).resolve()
        if target==root or root in target.parents:
            z.extract(member,root)

def _ingest_tree(root,subject,results,depth=0):
    for p in Path(root).rglob('*'):
        if not p.is_file():continue
        ext=p.suffix.lower()
        if ext in SUPPORTED:
            try:results.append(embed_and_store(p.name,subject,p))
            except Exception as e:results.append({'filename':p.name,'error':str(e),'passages':0,'embeddings':0})
        elif ext=='.zip' and depth<MAX_ZIP_DEPTH:
            try:
                with zipfile.ZipFile(p) as z:
                    child=p.parent/(p.stem+'_expanded')
                    child.mkdir(exist_ok=True)
                    _safe_extract(z,child)
                _ingest_tree(child,subject,results,depth+1)
            except Exception as e:results.append({'filename':p.name,'error':f'ZIP extraction failed: {e}','passages':0,'embeddings':0})

def ingest_path(path,subject):
    path=Path(path);results=[]
    if path.suffix.lower()=='.zip':
        with tempfile.TemporaryDirectory() as td:
            try:
                with zipfile.ZipFile(path) as z:_safe_extract(z,td)
            except zipfile.BadZipFile:
                return [{'filename':path.name,'error':'The uploaded ZIP is invalid or incomplete. Please re-create the ZIP and upload it again.','passages':0,'embeddings':0}]
            except Exception as e:
                return [{'filename':path.name,'error':f'ZIP extraction failed: {e}','passages':0,'embeddings':0}]
            _ingest_tree(td,subject,results)
    elif path.suffix.lower() in SUPPORTED:
        try:results.append(embed_and_store(path.name,subject,path))
        except Exception as e:results.append({'filename':path.name,'error':str(e),'passages':0,'embeddings':0})
    return results
